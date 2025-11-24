from pathlib import Path

import openpyxl
import pytest

from ade_engine.pipeline.io import iter_tables, list_input_files, read_table, sheet_name


def test_list_input_files_filters_extensions(tmp_path: Path) -> None:
    files = [
        tmp_path / "data.csv",
        tmp_path / "notes.txt",
        tmp_path / "table.xlsx",
    ]
    (tmp_path / "dir").mkdir()
    for path in files:
        if path.suffix == ".xlsx":
            workbook = openpyxl.Workbook()
            workbook.save(path)
            workbook.close()
        else:
            path.write_text("header\nvalue\n", encoding="utf-8")
    results = list_input_files(tmp_path)
    assert [file.name for file in results] == ["data.csv", "table.xlsx"]


def test_read_table_handles_csv(tmp_path: Path) -> None:
    csv_path = tmp_path / "values.csv"
    csv_path.write_text("Name,Email\nAlice,alice@example.com\n", encoding="utf-8")
    header, rows = read_table(csv_path)
    assert header == ["Name", "Email"]
    assert rows[0][1] == "alice@example.com"


def test_read_table_uses_named_sheet_when_provided(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    workbook.active.title = "Summary"
    data = workbook.create_sheet("Data")
    data.append(["Name", "Role"])
    data.append(["Mina", "Analyst"])
    path = tmp_path / "people.xlsx"
    workbook.save(path)
    workbook.close()

    header, rows = read_table(path, sheet_name="Data")
    assert header == ["Name", "Role"]
    assert rows == [["Mina", "Analyst"]]


def test_read_table_errors_when_sheet_missing(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    path = tmp_path / "missing.xlsx"
    workbook.save(path)
    workbook.close()

    with pytest.raises(RuntimeError):
        read_table(path, sheet_name="Nope")


def test_sheet_name_sanitizes_input() -> None:
    assert sheet_name("Employee-List 2024") == "Employee List 2024"


def test_iter_tables_yields_all_worksheets_by_default(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    active = workbook.active
    active.title = "Summary"
    active.append(["Name", "Role"])
    active.append(["Mina", "Analyst"])

    detail = workbook.create_sheet("Detail")
    detail.append(["Name", "Score"])
    detail.append(["Mina", "90"])

    path = tmp_path / "people.xlsx"
    workbook.save(path)
    workbook.close()

    tables = list(iter_tables(path))
    assert [sheet for sheet, _, _ in tables] == ["Summary", "Detail"]


def test_iter_tables_can_limit_to_subset(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    workbook.active.title = "Primary"
    workbook.active.append(["Name"])
    workbook.active.append(["Keep"])

    other = workbook.create_sheet("Secondary")
    other.append(["Name"])
    other.append(["Skip"])

    path = tmp_path / "subset.xlsx"
    workbook.save(path)
    workbook.close()

    tables = list(iter_tables(path, sheet_names=["Secondary"]))
    assert [sheet for sheet, _, _ in tables] == ["Secondary"]
