from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import polars as pl
from openpyxl.utils import get_column_letter

from ade_engine.extensions.registry import Registry
from ade_engine.infrastructure.settings import Settings
from ade_engine.models.events import (
    MAX_CANDIDATES,
    Candidate,
    CellsCount,
    ColumnsCount,
    ColumnHeader,
    ColumnStructure,
    Counts,
    DataInfo,
    Evaluation,
    Execution,
    Failure,
    FieldsCount,
    FieldOccurrences,
    FieldSummary,
    Finding,
    HeaderInfo,
    Mapping,
    Outputs,
    OutputsNormalized,
    Region,
    RowsCount,
    RunCompletedPayloadV1,
    SheetLocator,
    SheetRef,
    SheetScan,
    SheetSummary,
    TableLocator,
    TableRef,
    TableStructure,
    TableSummary,
    UnmappedReason,
    Validation,
    WorkbookLocator,
    WorkbookRef,
    WorkbookSummary,
)
from ade_engine.models.run import RunError, RunStatus
from ade_engine.models.table import TableRegion, TableResult

_SEVERITY_ORDER: dict[str, int] = {"info": 0, "warning": 1, "error": 2}


def _rfc3339_utc(ts: datetime) -> str:
    if ts.tzinfo is None:
        ts = ts.replace(tzinfo=timezone.utc)
    return ts.astimezone(timezone.utc).isoformat(timespec="milliseconds").replace("+00:00", "Z")


def _is_empty_cell(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip() == ""
    return False


def _normalize_header(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    lowered = text.lower()
    out: list[str] = []
    prev_underscore = False
    for ch in lowered:
        if ch.isalnum():
            out.append(ch)
            prev_underscore = False
        else:
            if not prev_underscore:
                out.append("_")
                prev_underscore = True
    normalized = "".join(out).strip("_")
    return normalized or None


def _is_placeholder_header(raw: str | None) -> bool:
    if raw is None:
        return True
    text = raw.strip()
    if not text:
        return True
    lowered = text.lower()
    if lowered in {"n/a", "na", "none", "null"}:
        return True
    if lowered.startswith("unnamed"):
        return True
    return False


def _max_severity(issues_by_severity: dict[str, int]) -> str | None:
    if not issues_by_severity:
        return None
    best = None
    best_rank = -1
    for sev, count in issues_by_severity.items():
        if count <= 0:
            continue
        rank = _SEVERITY_ORDER.get(sev, -1)
        if rank > best_rank:
            best = sev
            best_rank = rank
    return best


@dataclass
class _SheetAccum:
    index: int
    name: str
    scan: dict[str, Any] | None = None
    tables: list[TableResult] = field(default_factory=list)


@dataclass
class _WorkbookAccum:
    index: int
    name: str
    sheets: dict[int, _SheetAccum] = field(default_factory=dict)


class RunCompletionReportBuilder:
    """Accumulates per-table facts and builds the engine.run.completed payload (v1)."""

    def __init__(self, *, input_file: Path, settings: Settings) -> None:
        self._settings = settings
        self._registry: Registry | None = None
        self._workbook = _WorkbookAccum(index=0, name=input_file.name)

    def set_registry(self, registry: Registry) -> None:
        self._registry = registry

    def record_sheet_scan(self, *, sheet_index: int, sheet_name: str, scan: dict[str, Any]) -> None:
        sheet = self._workbook.sheets.get(sheet_index)
        if sheet is None:
            sheet = _SheetAccum(index=sheet_index, name=sheet_name)
            self._workbook.sheets[sheet_index] = sheet
        sheet.scan = dict(scan)

    def record_table(self, table: TableResult) -> None:
        sheet_index = int(getattr(table, "sheet_index", 0) or 0)
        sheet_name = str(getattr(table, "sheet_name", "") or "")
        sheet = self._workbook.sheets.get(sheet_index)
        if sheet is None:
            sheet = _SheetAccum(index=sheet_index, name=sheet_name)
            self._workbook.sheets[sheet_index] = sheet
        sheet.tables.append(table)

    # ------------------------------------------------------------------
    # Build
    # ------------------------------------------------------------------
    def build(
        self,
        *,
        run_status: RunStatus,
        started_at: datetime,
        completed_at: datetime,
        error: RunError | None,
        output_path: Path | None,
        output_written: bool,
    ) -> RunCompletedPayloadV1:
        expected_fields = list(self._registry.fields.values()) if self._registry is not None else []

        workbook_summary = self._build_workbook_summary(
            workbook=self._workbook,
            expected_fields=expected_fields,
            output_path=output_path,
            output_written=output_written,
        )

        run_workbooks = [workbook_summary]
        run_counts, run_validation, run_fields = self._rollup_run(expected_fields, run_workbooks)

        execution = self._build_execution(
            run_status=run_status,
            started_at=started_at,
            completed_at=completed_at,
            error=error,
        )
        evaluation = self._grade_run(execution=execution, counts=run_counts, validation=run_validation)

        outputs = None
        if output_written and output_path is not None:
            outputs = Outputs(normalized=OutputsNormalized(path=str(output_path)))

        return RunCompletedPayloadV1(
            execution=execution,
            evaluation=evaluation,
            counts=run_counts,
            validation=run_validation,
            fields=run_fields,
            outputs=outputs,
            workbooks=run_workbooks,
        )

    # ------------------------------------------------------------------
    # Node builders
    # ------------------------------------------------------------------
    def _build_execution(
        self,
        *,
        run_status: RunStatus,
        started_at: datetime,
        completed_at: datetime,
        error: RunError | None,
    ) -> Execution:
        duration_ms = int(max(0, round((completed_at - started_at).total_seconds() * 1000)))
        status = "succeeded" if run_status == RunStatus.SUCCEEDED else "failed"
        failure = None
        if status == "failed":
            stage = (error.stage or "unknown") if error is not None else "unknown"
            code = (error.code.value if error is not None else "unknown_error")
            message = (error.message if error is not None else "run failed")
            failure = Failure(stage=str(stage), code=str(code), message=str(message))

        return Execution(
            status=status,  # type: ignore[arg-type]
            started_at=_rfc3339_utc(started_at),
            completed_at=_rfc3339_utc(completed_at),
            duration_ms=duration_ms,
            failure=failure,
        )

    def _build_workbook_summary(
        self,
        *,
        workbook: _WorkbookAccum,
        expected_fields: list[Any],
        output_path: Path | None,
        output_written: bool,
    ) -> WorkbookSummary:
        wb_ref = WorkbookRef(index=workbook.index, name=workbook.name)
        sheets: list[SheetSummary] = []
        for sheet_index in sorted(workbook.sheets):
            sheet = workbook.sheets[sheet_index]
            sheets.append(
                self._build_sheet_summary(
                    workbook_ref=wb_ref,
                    sheet=sheet,
                    expected_fields=expected_fields,
                    output_path=output_path,
                    output_written=output_written,
                )
            )

        counts, validation, fields = self._rollup_workbook(expected_fields, sheets)
        return WorkbookSummary(
            locator=WorkbookLocator(workbook=wb_ref),
            counts=counts,
            validation=validation,
            fields=fields,
            sheets=sheets,
        )

    def _build_sheet_summary(
        self,
        *,
        workbook_ref: WorkbookRef,
        sheet: _SheetAccum,
        expected_fields: list[Any],
        output_path: Path | None,
        output_written: bool,
    ) -> SheetSummary:
        sheet_ref = SheetRef(index=sheet.index, name=sheet.name)
        tables: list[TableSummary] = []
        for t in sorted(sheet.tables, key=lambda x: int(getattr(x, "table_index", 0) or 0)):
            tables.append(
                self._build_table_summary(
                    workbook_ref=workbook_ref,
                    sheet_ref=sheet_ref,
                    table=t,
                    expected_fields=expected_fields,
                    output_path=output_path,
                    output_written=output_written,
                )
            )

        counts, validation, fields = self._rollup_sheet(expected_fields, tables)

        scan = None
        if sheet.scan is not None:
            scan = SheetScan(
                rows_emitted=int(sheet.scan.get("rows_emitted", 0) or 0),
                stopped_early=bool(sheet.scan.get("stopped_early", False)),
                truncated_rows=int(sheet.scan.get("truncated_rows", 0) or 0),
            )

        return SheetSummary(
            locator=SheetLocator(workbook=workbook_ref, sheet=sheet_ref),
            counts=counts,
            validation=validation,
            fields=fields,
            scan=scan,
            tables=tables,
        )

    def _build_table_summary(
        self,
        *,
        workbook_ref: WorkbookRef,
        sheet_ref: SheetRef,
        table: TableResult,
        expected_fields: list[Any],
        output_path: Path | None,
        output_written: bool,
    ) -> TableSummary:
        table_index = int(getattr(table, "table_index", 0) or 0)
        row_count = int(getattr(table, "row_count", 0) or 0)
        source_cols = list(getattr(table, "source_columns", []) or [])
        col_total = len(source_cols)

        source_region = getattr(table, "source_region", None)
        if isinstance(source_region, TableRegion):
            header_row_start = source_region.header_row
            data_row_start = source_region.data_first_row
            data_row_count = source_region.data_row_count
            header_inferred = False
            region_a1 = source_region.a1
        else:
            header_row_start = 1
            data_row_start = 2
            data_row_count = row_count
            header_inferred = False

            data_end_row_number = data_row_start + data_row_count - 1 if data_row_count > 0 else header_row_start
            col_letter = get_column_letter(max(1, col_total))
            region_a1 = f"A{header_row_start}:{col_letter}{data_end_row_number}"

        mapped_by_index = {int(c.source_index): c for c in (getattr(table, "mapped_columns", []) or [])}
        scores_by_column: dict[int, dict[str, float]] = dict(getattr(table, "column_scores", {}) or {})
        duplicate_unmapped: set[int] = set(getattr(table, "duplicate_unmapped_indices", set()) or set())

        empty_cols = 0
        non_empty_cells_total = 0
        columns: list[ColumnStructure] = []

        mapped_count = 0
        ambiguous_count = 0
        unmapped_count = 0
        passthrough_count = 0

        for col in source_cols:
            raw_header = None if col.header in (None, "") else str(col.header)
            if raw_header is not None:
                raw_header = raw_header.strip()
                if raw_header == "":
                    raw_header = None
            normalized_header = _normalize_header(raw_header) if raw_header is not None else None
            header = ColumnHeader(raw=raw_header, normalized=normalized_header)

            values = list(getattr(col, "values", []) or [])
            non_empty_cells = sum(1 for v in values if not _is_empty_cell(v))
            if non_empty_cells == 0:
                empty_cols += 1
            non_empty_cells_total += non_empty_cells

            mapping, status_bucket = self._build_column_mapping(
                col_index=int(col.index),
                header_raw=raw_header,
                mapped=mapped_by_index.get(int(col.index)),
                scores=scores_by_column.get(int(col.index), {}),
                duplicate_unmapped=duplicate_unmapped,
            )

            if status_bucket == "mapped":
                mapped_count += 1
            elif status_bucket == "ambiguous":
                ambiguous_count += 1
            elif status_bucket == "unmapped":
                unmapped_count += 1
            else:
                passthrough_count += 1

            columns.append(
                ColumnStructure(
                    index=int(col.index),
                    header=header,
                    non_empty_cells=int(non_empty_cells),
                    mapping=mapping,
                )
            )

        # TableResult.row_count reflects the post-hook output table height, but
        # SourceColumn.values reflect the detected source region. Hooks may
        # filter rows (e.g. drop invalid records), so compute structural counts
        # against the source row count for internal consistency.
        source_row_count = data_row_count
        if source_cols:
            source_row_count = max(
                source_row_count,
                max(len(getattr(col, "values", []) or []) for col in source_cols),
            )

        empty_rows = self._count_empty_rows(source_cols, data_row_count=source_row_count)

        counts = Counts(
            rows=RowsCount(total=source_row_count, empty=empty_rows),
            columns=ColumnsCount(
                total=col_total,
                empty=empty_cols,
                mapped=mapped_count,
                ambiguous=ambiguous_count,
                unmapped=unmapped_count,
                passthrough=passthrough_count,
            ),
            fields=FieldsCount(expected=len(expected_fields), mapped=len(self._mapped_fields_in_table(table, expected_fields))),
            cells=CellsCount(total=source_row_count * col_total, non_empty=non_empty_cells_total)
            if col_total and source_row_count
            else CellsCount(total=0, non_empty=0),
        )

        validation = self._build_validation(table)

        structure = TableStructure(
            region=Region(a1=region_a1),
            header=HeaderInfo(row_start=header_row_start, row_count=1, inferred=header_inferred),
            data=DataInfo(row_start=data_row_start, row_count=data_row_count),
            columns=sorted(columns, key=lambda c: c.index),
        )

        outputs = None
        if output_written and output_path is not None:
            output_region = getattr(table, "output_region", None)
            out_sheet_name = getattr(table, "output_sheet_name", None)
            if (
                isinstance(output_region, TableRegion)
                and isinstance(out_sheet_name, str)
                and out_sheet_name.strip()
            ):
                outputs = Outputs(
                    normalized=OutputsNormalized(
                        path=str(output_path),
                        sheet_name=str(getattr(table, "sheet_name", "") or ""),
                        range_a1=f"{out_sheet_name}!{output_region.a1}",
                    )
                )

        return TableSummary(
            locator=TableLocator(
                workbook=workbook_ref,
                sheet=sheet_ref,
                table=TableRef(index=table_index),
            ),
            counts=counts,
            validation=validation,
            structure=structure,
            outputs=outputs,
        )

    def _build_column_mapping(
        self,
        *,
        col_index: int,
        header_raw: str | None,
        mapped: Any | None,
        scores: dict[str, float],
        duplicate_unmapped: set[int],
    ) -> tuple[Mapping, str]:
        candidates = self._top_candidates(scores)

        if mapped is not None:
            field_name = str(getattr(mapped, "field_name", "") or "")
            score_val = getattr(mapped, "score", None)
            if score_val is None:
                score_val = next((c.score for c in candidates if c.field == field_name), 0.0)
            return (
                Mapping(
                    status="mapped",
                    field=field_name,
                    score=float(max(0.0, float(score_val))),
                    method="classifier",
                    candidates=candidates,
                ),
                "mapped",
            )

        if _is_placeholder_header(header_raw):
            return (
                Mapping(status="unmapped", unmapped_reason="empty_or_placeholder_header"),
                "unmapped",
            )

        if col_index in duplicate_unmapped:
            return (
                Mapping(status="unmapped", candidates=candidates, unmapped_reason="duplicate_field"),
                "unmapped",
            )

        # v1 default: columns without a selected field may be carried through as raw output.
        if not self._settings.remove_unmapped_columns:
            return (
                Mapping(status="passthrough", candidates=candidates, unmapped_reason="passthrough_policy"),
                "passthrough",
            )

        # If passthrough is disabled, include reason codes for analysis.
        reason: UnmappedReason = "below_threshold" if candidates else "no_signal"
        return (
            Mapping(status="unmapped", candidates=candidates, unmapped_reason=reason),
            "unmapped",
        )

    def _top_candidates(self, scores: dict[str, float]) -> list[Candidate]:
        items: list[tuple[str, float]] = []
        for field, score in (scores or {}).items():
            try:
                s = float(score)
            except Exception:
                continue
            if s <= 0:
                continue
            items.append((str(field), s))
        items.sort(key=lambda kv: (-kv[1], kv[0]))
        items = items[:MAX_CANDIDATES]
        return [Candidate(field=f, score=float(s)) for f, s in items]

    def _count_empty_rows(self, source_cols: list[Any], *, data_row_count: int) -> int:
        if data_row_count <= 0:
            return 0
        empty = 0
        for row_idx in range(data_row_count):
            any_non_empty = False
            for col in source_cols:
                vals = getattr(col, "values", None)
                if not isinstance(vals, list):
                    continue
                if row_idx < len(vals) and not _is_empty_cell(vals[row_idx]):
                    any_non_empty = True
                    break
            if not any_non_empty:
                empty += 1
        return empty

    def _build_validation(self, table: TableResult) -> Validation:
        df = getattr(table, "table", None)
        issues_total = 0
        if isinstance(df, pl.DataFrame) and df.height:
            if "__ade_issue_count" in df.columns:
                raw_total = df.get_column("__ade_issue_count").sum()
                issues_total = int(raw_total or 0)
            else:
                issue_cols = [c for c in df.columns if c.startswith("__ade_issue__")]
                for col in issue_cols:
                    raw_count = df.get_column(col).is_not_null().sum()
                    issues_total += int(raw_count or 0)

        by_sev: dict[str, int] = {"warning": issues_total} if issues_total else {}
        return Validation(
            rows_evaluated=int(getattr(table, "row_count", 0) or 0),
            issues_total=int(issues_total),
            issues_by_severity={k: int(v) for k, v in by_sev.items()} if issues_total else {},
            max_severity=_max_severity(by_sev),
        )

    # ------------------------------------------------------------------
    # Rollups
    # ------------------------------------------------------------------
    def _expected_field_names(self, expected_fields: list[Any]) -> list[str]:
        return [str(getattr(f, "name", "") or "") for f in expected_fields]

    def _mapped_fields_in_table(self, table: TableResult, expected_fields: list[Any]) -> set[str]:
        expected = set(self._expected_field_names(expected_fields))
        mapped = {str(getattr(c, "field_name", "") or "") for c in (getattr(table, "mapped_columns", []) or [])}
        return {f for f in mapped if f in expected}

    def _rollup_sheet(self, expected_fields: list[Any], tables: list[TableSummary]) -> tuple[Counts, Validation, list[FieldSummary]]:
        rows_total = rows_empty = 0
        cols_total = cols_empty = 0
        cols_mapped = cols_ambiguous = cols_unmapped = cols_passthrough = 0
        cells_total = cells_non_empty = 0

        validation = self._zero_validation()
        field_occ = self._field_occurrences_init(expected_fields)

        for t in tables:
            rows_total += t.counts.rows.total
            rows_empty += t.counts.rows.empty
            cols_total += t.counts.columns.total
            cols_empty += t.counts.columns.empty
            cols_mapped += t.counts.columns.mapped
            cols_ambiguous += t.counts.columns.ambiguous
            cols_unmapped += t.counts.columns.unmapped
            cols_passthrough += t.counts.columns.passthrough
            if t.counts.cells is not None:
                cells_total += t.counts.cells.total
                cells_non_empty += t.counts.cells.non_empty

            validation.rows_evaluated += t.validation.rows_evaluated
            validation.issues_total += t.validation.issues_total
            for sev, count in t.validation.issues_by_severity.items():
                validation.issues_by_severity[sev] = validation.issues_by_severity.get(sev, 0) + count

        self._accumulate_field_occurrences(field_occ, expected_fields, tables)
        validation.max_severity = _max_severity(validation.issues_by_severity)
        fields = self._field_summaries(expected_fields, field_occ)

        counts = Counts(
            tables=len(tables),
            rows=RowsCount(total=rows_total, empty=rows_empty),
            columns=ColumnsCount(
                total=cols_total,
                empty=cols_empty,
                mapped=cols_mapped,
                ambiguous=cols_ambiguous,
                unmapped=cols_unmapped,
                passthrough=cols_passthrough,
            ),
            fields=FieldsCount(expected=len(expected_fields), mapped=sum(1 for f in fields if f.mapped)),
            cells=CellsCount(total=cells_total, non_empty=cells_non_empty),
        )
        return counts, validation, fields

    def _rollup_workbook(self, expected_fields: list[Any], sheets: list[SheetSummary]) -> tuple[Counts, Validation, list[FieldSummary]]:
        rows_total = rows_empty = 0
        cols_total = cols_empty = 0
        cols_mapped = cols_ambiguous = cols_unmapped = cols_passthrough = 0
        cells_total = cells_non_empty = 0
        tables_total = 0

        validation = self._zero_validation()
        field_occ = self._field_occurrences_init(expected_fields)

        for s in sheets:
            tables_total += int(s.counts.tables or 0)
            rows_total += s.counts.rows.total
            rows_empty += s.counts.rows.empty
            cols_total += s.counts.columns.total
            cols_empty += s.counts.columns.empty
            cols_mapped += s.counts.columns.mapped
            cols_ambiguous += s.counts.columns.ambiguous
            cols_unmapped += s.counts.columns.unmapped
            cols_passthrough += s.counts.columns.passthrough
            if s.counts.cells is not None:
                cells_total += s.counts.cells.total
                cells_non_empty += s.counts.cells.non_empty

            validation.rows_evaluated += s.validation.rows_evaluated
            validation.issues_total += s.validation.issues_total
            for sev, count in s.validation.issues_by_severity.items():
                validation.issues_by_severity[sev] = validation.issues_by_severity.get(sev, 0) + count

            for f in s.fields:
                occ = field_occ.get(f.field)
                if occ is None or not f.mapped:
                    continue
                occ["tables"] += f.occurrences.tables
                occ["columns"] += f.occurrences.columns
                occ["best"] = max(occ["best"], float(f.best_mapping_score or 0.0))

        validation.max_severity = _max_severity(validation.issues_by_severity)
        fields = self._field_summaries(expected_fields, field_occ)

        counts = Counts(
            sheets=len(sheets),
            tables=tables_total,
            rows=RowsCount(total=rows_total, empty=rows_empty),
            columns=ColumnsCount(
                total=cols_total,
                empty=cols_empty,
                mapped=cols_mapped,
                ambiguous=cols_ambiguous,
                unmapped=cols_unmapped,
                passthrough=cols_passthrough,
            ),
            fields=FieldsCount(expected=len(expected_fields), mapped=sum(1 for f in fields if f.mapped)),
            cells=CellsCount(total=cells_total, non_empty=cells_non_empty),
        )
        return counts, validation, fields

    def _rollup_run(self, expected_fields: list[Any], workbooks: list[WorkbookSummary]) -> tuple[Counts, Validation, list[FieldSummary]]:
        rows_total = rows_empty = 0
        cols_total = cols_empty = 0
        cols_mapped = cols_ambiguous = cols_unmapped = cols_passthrough = 0
        cells_total = cells_non_empty = 0
        sheets_total = 0
        tables_total = 0

        validation = self._zero_validation()
        field_occ = self._field_occurrences_init(expected_fields)

        for w in workbooks:
            sheets_total += int(w.counts.sheets or 0)
            tables_total += int(w.counts.tables or 0)
            rows_total += w.counts.rows.total
            rows_empty += w.counts.rows.empty
            cols_total += w.counts.columns.total
            cols_empty += w.counts.columns.empty
            cols_mapped += w.counts.columns.mapped
            cols_ambiguous += w.counts.columns.ambiguous
            cols_unmapped += w.counts.columns.unmapped
            cols_passthrough += w.counts.columns.passthrough
            if w.counts.cells is not None:
                cells_total += w.counts.cells.total
                cells_non_empty += w.counts.cells.non_empty

            validation.rows_evaluated += w.validation.rows_evaluated
            validation.issues_total += w.validation.issues_total
            for sev, count in w.validation.issues_by_severity.items():
                validation.issues_by_severity[sev] = validation.issues_by_severity.get(sev, 0) + count

            for f in w.fields:
                occ = field_occ.get(f.field)
                if occ is None or not f.mapped:
                    continue
                occ["tables"] += f.occurrences.tables
                occ["columns"] += f.occurrences.columns
                occ["best"] = max(occ["best"], float(f.best_mapping_score or 0.0))

        validation.max_severity = _max_severity(validation.issues_by_severity)
        fields = self._field_summaries(expected_fields, field_occ)

        counts = Counts(
            workbooks=len(workbooks),
            sheets=sheets_total,
            tables=tables_total,
            rows=RowsCount(total=rows_total, empty=rows_empty),
            columns=ColumnsCount(
                total=cols_total,
                empty=cols_empty,
                mapped=cols_mapped,
                ambiguous=cols_ambiguous,
                unmapped=cols_unmapped,
                passthrough=cols_passthrough,
            ),
            fields=FieldsCount(expected=len(expected_fields), mapped=sum(1 for f in fields if f.mapped)),
            cells=CellsCount(total=cells_total, non_empty=cells_non_empty),
        )
        return counts, validation, fields

    def _zero_validation(self) -> Validation:
        return Validation(rows_evaluated=0, issues_total=0, issues_by_severity={}, max_severity=None)

    def _field_occurrences_init(self, expected_fields: list[Any]) -> dict[str, dict[str, Any]]:
        out: dict[str, dict[str, Any]] = {}
        for f in expected_fields:
            name = str(getattr(f, "name", "") or "")
            out[name] = {"tables": 0, "columns": 0, "best": 0.0}
        return out

    def _accumulate_field_occurrences(
        self, occ: dict[str, dict[str, Any]], expected_fields: list[Any], tables: list[TableSummary]
    ) -> None:
        expected_names = set(self._expected_field_names(expected_fields))
        for t in tables:
            seen_in_table: set[str] = set()
            for col in t.structure.columns:
                m = col.mapping
                if m.status != "mapped" or m.field is None:
                    continue
                if m.field not in expected_names:
                    continue
                record = occ.get(m.field)
                if record is None:
                    continue
                record["columns"] += 1
                record["best"] = max(record["best"], float(m.score or 0.0))
                seen_in_table.add(m.field)
            for field in seen_in_table:
                occ[field]["tables"] += 1

    def _field_summaries(self, expected_fields: list[Any], occ: dict[str, dict[str, Any]]) -> list[FieldSummary]:
        out: list[FieldSummary] = []
        for f in expected_fields:
            name = str(getattr(f, "name", "") or "")
            label = getattr(f, "label", None)
            rec = occ.get(name) or {"tables": 0, "columns": 0, "best": 0.0}
            mapped = bool(rec["tables"] > 0)
            out.append(
                FieldSummary(
                    field=name,
                    label=str(label) if label is not None else None,
                    mapped=mapped,
                    best_mapping_score=float(rec["best"]) if mapped else None,
                    occurrences=FieldOccurrences(tables=int(rec["tables"]), columns=int(rec["columns"])),
                )
            )
        return out

    # ------------------------------------------------------------------
    # Evaluation grading
    # ------------------------------------------------------------------
    def _grade_run(self, *, execution: Execution, counts: Counts, validation: Validation) -> Evaluation:
        findings: list[Finding] = []

        tables = int(counts.tables or 0)
        expected = int(counts.fields.expected)
        mapped_fields = int(counts.fields.mapped)

        if execution.status == "failed" and tables == 0:
            findings.append(Finding(code="execution_failed", severity="error", message="Execution failed"))
            return Evaluation(outcome="unknown", findings=findings)

        if tables == 0:
            findings.append(Finding(code="no_tables_detected", severity="error", message="No tables detected"))
            outcome: str = "failure"
        elif mapped_fields == 0 and expected > 0:
            findings.append(Finding(code="no_fields_mapped", severity="error", message="No expected fields mapped"))
            outcome = "failure"
        elif expected > 0 and mapped_fields < expected:
            findings.append(
                Finding(
                    code="fields_unmapped",
                    severity="warning",
                    message=f"Only {mapped_fields}/{expected} expected fields were mapped",
                )
            )
            outcome = "partial"
        else:
            outcome = "success"

        warn_count = int(validation.issues_by_severity.get("warning", 0) or 0)
        err_count = int(validation.issues_by_severity.get("error", 0) or 0)
        if warn_count > 0:
            findings.append(
                Finding(
                    code="validation_warnings_present",
                    severity="warning",
                    message="Validation warnings were emitted",
                    count=warn_count,
                )
            )
        if err_count > 0:
            findings.append(
                Finding(
                    code="validation_errors_present",
                    severity="error",
                    message="Validation errors were emitted",
                    count=err_count,
                )
            )
            if outcome == "success":
                outcome = "partial"

        if execution.status == "failed" and tables > 0:
            findings.append(Finding(code="execution_failed", severity="error", message="Execution failed"))
            if outcome == "success":
                outcome = "partial"

        return Evaluation(outcome=outcome, findings=findings)  # type: ignore[arg-type]
