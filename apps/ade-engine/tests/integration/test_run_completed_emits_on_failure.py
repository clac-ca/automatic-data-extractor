from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))

from ade_engine.application.engine import Engine
from ade_engine.infrastructure.settings import Settings
from ade_engine.models.events import RunCompletedPayloadV1
from ade_engine.models.run import RunRequest, RunStatus


def _write_config_package(root: Path) -> None:
    pkg = root / "ade_config"
    pkg.mkdir(parents=True, exist_ok=True)
    (pkg / "__init__.py").write_text(
        """
def register(registry):
    from . import rows, columns, hooks
    rows.register(registry)
    columns.register(registry)
    hooks.register(registry)
""",
        encoding="utf-8",
    )

    (pkg / "rows.py").write_text(
        """
from ade_engine.models import RowKind

def pick_header(*, row_values, **_):
    normalized = {str(v).strip().lower() for v in row_values or [] if v not in (None, "")}
    if {"email", "name"}.issubset(normalized):
        return {"header": 1.0}
    return {}

def register(registry):
    registry.register_row_detector(pick_header, row_kind=RowKind.HEADER.value, priority=10)
""",
        encoding="utf-8",
    )

    (pkg / "columns.py").write_text(
        """
from ade_engine.models import FieldDef

def detect_email(*, header, **_):
    header = (header or "").strip().lower()
    return {"email": 1.0} if header == "email" else None

def detect_name(*, header, **_):
    header = (header or "").strip().lower()
    return {"name": 1.0} if header == "name" else None

def register(registry):
    registry.register_field(FieldDef(name="email"))
    registry.register_field(FieldDef(name="name"))
    registry.register_column_detector(detect_email, field="email", priority=10)
    registry.register_column_detector(detect_name, field="name", priority=10)
""",
        encoding="utf-8",
    )

    (pkg / "hooks.py").write_text(
        """
from ade_engine.models import HookName

def register(registry):
    registry.register_hook(on_workbook_before_save, hook_name=HookName.ON_WORKBOOK_BEFORE_SAVE, priority=0)

def on_workbook_before_save(**_):
    raise RuntimeError("boom before save")
""",
        encoding="utf-8",
    )


def _write_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Email", "Name"])
    ws.append(["a@example.com", "Alice"])
    wb.save(path)


def test_run_completed_emitted_once_on_failure(tmp_path: Path) -> None:
    config_root = tmp_path / "cfg"
    config_root.mkdir()
    _write_config_package(config_root)

    source = tmp_path / "input.xlsx"
    _write_workbook(source)

    engine = Engine(settings=Settings(log_format="ndjson"))
    output_dir = tmp_path / "out"
    logs_dir = tmp_path / "logs"
    req = RunRequest(config_package=config_root, input_file=source, output_dir=output_dir, logs_dir=logs_dir)
    result = engine.run(req)
    assert result.status == RunStatus.FAILED

    log_path = logs_dir / f"{source.stem}_engine_events.ndjson"
    assert log_path.exists()

    completed_events: list[dict] = []
    for line in log_path.read_text(encoding="utf-8").splitlines():
        rec = json.loads(line)
        if rec.get("event") == "engine.run.completed":
            completed_events.append(rec)

    assert len(completed_events) == 1
    payload = completed_events[0].get("data")
    model = RunCompletedPayloadV1.model_validate(payload, strict=True)

    assert model.execution.status == "failed"
    assert model.counts.tables is not None and model.counts.tables >= 1
    assert model.outputs is None
