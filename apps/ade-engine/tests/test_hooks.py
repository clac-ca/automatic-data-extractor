from __future__ import annotations

import copy
import json
import sys
import textwrap
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import Workbook

from ade_engine.config.hook_registry import HookStage
from ade_engine.config.loader import load_config_runtime
from ade_engine.core.errors import HookError
from ade_engine.core.hooks import run_hooks
from ade_engine.core.types import ExtractedTable, RunContext, RunPaths, RunResult, RunStatus
from ade_engine.infra.telemetry import PipelineLogger

BASE_MANIFEST = {
    "schema": "ade.manifest/v1",
    "version": "1.0.0",
    "name": "Hook Config",
    "description": None,
    "script_api_version": 2,
    "columns": {"order": [], "fields": {}},
    "hooks": {
        "on_run_start": [],
        "on_after_extract": [],
        "on_after_mapping": [],
        "on_before_save": [],
        "on_run_end": [],
    },
    "writer": {
        "append_unmapped_columns": True,
        "unmapped_prefix": "raw_",
        "output_sheet": "Normalized",
    },
    "extra": None,
}


def _purge_config_modules(prefix: str = "ade_config") -> None:
    for name in list(sys.modules):
        if name == prefix or name.startswith(f"{prefix}."):
            sys.modules.pop(name)


@pytest.fixture
def config_package(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    _purge_config_modules()
    pkg_root = tmp_path / "ade_config"
    pkg_root.mkdir()
    (pkg_root / "__init__.py").write_text("")
    monkeypatch.syspath_prepend(str(tmp_path))
    yield pkg_root
    _purge_config_modules()


def _write_manifest(pkg_root: Path, *, hooks: dict[str, list[str]]) -> Path:
    manifest = copy.deepcopy(BASE_MANIFEST)
    manifest["hooks"].update(hooks)
    manifest_path = pkg_root / "manifest.json"
    manifest_path.write_text(json.dumps(manifest))
    return manifest_path


def _write_hook(pkg_root: Path, name: str, body: str) -> None:
    hooks_dir = pkg_root / "hooks"
    hooks_dir.mkdir(exist_ok=True)
    (hooks_dir / "__init__.py").write_text("", encoding="utf-8")
    (hooks_dir / f"{name}.py").write_text(textwrap.dedent(body), encoding="utf-8")


def _run_and_logger(manifest_context, tmp_path: Path) -> tuple[RunContext, PipelineLogger]:
    paths = RunPaths(
        input_dir=tmp_path / "input",
        output_dir=tmp_path / "output",
        logs_dir=tmp_path / "logs",
    )
    for path in (paths.input_dir, paths.output_dir, paths.logs_dir):
        path.mkdir()

    run = RunContext(
        run_id=uuid4(),
        metadata={},
        manifest=manifest_context,
        paths=paths,
        started_at=datetime.now(timezone.utc),
    )
    return run, PipelineLogger(run=run)


def test_after_extract_hooks_apply_in_order_and_propagate_tables(
    config_package: Path, tmp_path: Path
) -> None:
    manifest_path = _write_manifest(
        config_package, hooks={"on_after_extract": ["hooks.first", "hooks.second"]}
    )
    _write_hook(
        config_package,
        "first",
        """
from typing import Sequence
from ade_engine.core.types import ExtractedTable

def run(*, tables: Sequence[ExtractedTable] | None, run, logger, **_):
    run.state.setdefault("order", []).append("first")
    logger.note("first hook")
    return list(tables or [])[:1]
""",
    )
    _write_hook(
        config_package,
        "second",
        """
def run(*, tables, run, **_):
    run.state.setdefault("order", []).append("second")
    if tables:
        tables[0].header_row.append("extra")
    return tables
""",
    )

    runtime = load_config_runtime(package="ade_config", manifest_path=manifest_path)
    run, logger = _run_and_logger(runtime.manifest, tmp_path)

    tables = [
        ExtractedTable(
            source_file=run.paths.input_dir / "data.csv",
            source_sheet=None,
            table_index=1,
            header_row=["col1"],
            data_rows=[["a"], ["b"]],
            header_row_index=1,
            first_data_row_index=2,
            last_data_row_index=3,
        ),
        ExtractedTable(
            source_file=run.paths.input_dir / "other.csv",
            source_sheet=None,
            table_index=2,
            header_row=["ignored"],
            data_rows=[["c"]],
            header_row_index=1,
            first_data_row_index=2,
            last_data_row_index=2,
        ),
    ]

    context = run_hooks(
        stage=HookStage.ON_AFTER_EXTRACT,
        registry=runtime.hooks,
        run=run,
        manifest=runtime.manifest,
        logger=logger,
        tables=tables,
    )

    assert run.state["order"] == ["first", "second"]
    assert context.tables is not None and len(context.tables) == 1
    assert context.tables[0].header_row == ["col1", "extra"]


def test_on_before_save_prefers_workbook_returned_by_hook(config_package: Path, tmp_path: Path) -> None:
    manifest_path = _write_manifest(config_package, hooks={"on_before_save": ["hooks.save"]})
    _write_hook(
        config_package,
        "save",
        """
from openpyxl import Workbook

def run(*, workbook, logger, **_):
    logger.note("replacing workbook")
    wb = Workbook()
    sheet = wb.active
    sheet.title = "custom"
    sheet["A1"] = "ok"
    return wb
""",
    )

    runtime = load_config_runtime(package="ade_config", manifest_path=manifest_path)
    run, logger = _run_and_logger(runtime.manifest, tmp_path)

    original = Workbook()
    original.active.title = "original"

    context = run_hooks(
        stage=HookStage.ON_BEFORE_SAVE,
        registry=runtime.hooks,
        run=run,
        manifest=runtime.manifest,
        logger=logger,
        workbook=original,
    )

    assert context.workbook is not original
    assert context.workbook is not None
    assert context.workbook.active.title == "custom"
    assert context.workbook.active["A1"].value == "ok"


def test_hook_error_includes_stage_and_module_path(config_package: Path, tmp_path: Path) -> None:
    manifest_path = _write_manifest(config_package, hooks={"on_run_end": ["hooks.fail"]})
    _write_hook(
        config_package,
        "fail",
        """
def run(*_, **__):
    raise RuntimeError("boom")
""",
    )

    runtime = load_config_runtime(package="ade_config", manifest_path=manifest_path)
    run, logger = _run_and_logger(runtime.manifest, tmp_path)
    result = RunResult(
        status=RunStatus.SUCCEEDED,
        error=None,
        run_id=run.run_id,
        output_paths=(),
        logs_dir=run.paths.logs_dir,
        processed_files=(),
    )

    with pytest.raises(HookError) as excinfo:
        run_hooks(
            stage=HookStage.ON_RUN_END,
            registry=runtime.hooks,
            run=run,
            manifest=runtime.manifest,
            logger=logger,
            result=result,
        )

    message = str(excinfo.value)
    assert "on_run_end" in message
    assert "ade_config.hooks.fail" in message
