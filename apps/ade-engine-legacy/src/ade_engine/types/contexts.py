"""Context objects passed into hooks and pipeline stages."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.worksheet import Worksheet

from ade_engine.config.manifest import ManifestContext
from ade_engine.types.mapping import ColumnMappingPatch
from ade_engine.types.origin import TableOrigin, TablePlacement, TableRegion
from ade_engine.types.tables import ExtractedTable, MappedTable, NormalizedTable


@dataclass
class RunContext:
    source_path: Path
    output_path: Path
    manifest: ManifestContext
    source_workbook: openpyxl.Workbook
    output_workbook: openpyxl.Workbook
    state: dict[str, Any] = field(default_factory=dict)
    logger: Any | None = None


@dataclass
class WorksheetContext:
    run: RunContext
    sheet_index: int
    source_worksheet: Worksheet
    output_worksheet: Worksheet


@dataclass
class TableContext:
    sheet: WorksheetContext
    origin: TableOrigin
    region: TableRegion
    extracted: ExtractedTable | None = None
    mapped: MappedTable | None = None
    normalized: NormalizedTable | None = None
    placement: TablePlacement | None = None
    view: "TableView" | None = None
    mapping_patch: ColumnMappingPatch | None = None


@dataclass(frozen=True)
class TableView:
    worksheet: Worksheet
    cell_range: CellRange

    @property
    def bounds(self) -> tuple[int, int, int, int]:
        return self.cell_range.bounds

    def header_cells(self):
        min_col, min_row, max_col, _ = self.bounds
        return [self.worksheet.cell(row=min_row, column=col) for col in range(min_col, max_col + 1)]

    def iter_data_rows(self):
        min_col, min_row, max_col, max_row = self.bounds
        for row_idx in range(min_row + 1, max_row + 1):
            yield [self.worksheet.cell(row=row_idx, column=col) for col in range(min_col, max_col + 1)]


__all__ = ["RunContext", "WorksheetContext", "TableContext", "TableView"]
